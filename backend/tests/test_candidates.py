"""
Unit tests for resume parser and candidate onboarding CRUD APIs.
"""

import io
from collections.abc import AsyncGenerator

import docx
import pytest
from httpx import ASGITransport, AsyncClient
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.auth import create_access_token
from app.core.database import get_db
from app.ingestion.resume_parser import (
    detect_is_linkedin_pdf,
    extract_metrics,
    extract_skills_from_text,
    extract_text_from_docx,
    parse_linkedin_text,
    parse_resume_bytes,
)
from app.main import app


def test_extract_skills_from_text() -> None:
    """Verify canonical technical skills normalization via SKILL_TAXONOMY."""
    sample_text = (
        "Built distributed event queues in Python using Kafka and PostgreSQL. "
        "Deployed on Kubernetes cluster with FastAPI and Docker."
    )
    skills = extract_skills_from_text(sample_text)
    assert "Python" in skills
    assert "Kafka" in skills
    assert "PostgreSQL" in skills
    assert "Kubernetes" in skills
    assert "FastAPI" in skills
    assert "Docker" in skills


def test_extract_metrics() -> None:
    """Verify metric detection for latency, percentages, scale, and hackathons."""
    text = (
        "Awarded 1st place overall at ETHGlobal Hackathon. "
        "Reduced p99 query latency by 45% and cut costs by $250k. "
        "Scaled distributed engine to 2M req/s across 100,000+ active users."
    )
    metrics = extract_metrics(text)
    assert "percentage_impact" in metrics
    assert "financial_impact" in metrics
    assert "throughput_or_latency" in metrics
    assert "user_or_data_scale" in metrics
    assert "honor_placement" in metrics
    assert "1st place" in metrics["honor_placement"].lower()


def test_detect_and_parse_linkedin_text() -> None:
    """Verify LinkedIn 'Save to PDF' detection and structured parsing."""
    sample_linkedin = (
        "Contact\n"
        "jane.developer@example.com\n"
        "www.linkedin.com/in/jane-dev\n\n"
        "Top Skills\n"
        "Python\n"
        "FastAPI\n"
        "Apache Kafka\n\n"
        "Languages\n"
        "English\n\n"
        "Jane Developer\n"
        "Senior Backend & AI Infrastructure Engineer\n"
        "San Francisco Bay Area\n\n"
        "Summary\n"
        "Distributed systems engineer with 6 years experience in low-latency "
        "systems.\n\n"
        "Experience\n"
        "Senior Systems Engineer\n"
        "CloudScale Inc\n"
        "Designed asynchronous event-driven data pipeline using Kafka and Go. "
        "Decreased system latency by 40% and handled 1M daily requests.\n\n"
        "1st Place Winner - AI Track\n"
        "HackMIT\n"
        "Built decentralized model serving engine with vLLM and Redis in 48 hours."
    )
    assert detect_is_linkedin_pdf(sample_linkedin) is True
    parsed = parse_linkedin_text(sample_linkedin)

    assert parsed.full_name == "Jane Developer"
    assert parsed.email == "jane.developer@example.com"
    assert parsed.linkedin_url == "https://www.linkedin.com/in/jane-dev"
    assert len(parsed.blocks) >= 2

    # Check hackathon category detection
    hack_block = next(
        (b for b in parsed.blocks if b.category == "hackathon_achievement"), None
    )
    assert hack_block is not None
    assert hack_block.is_quantified is True


def test_extract_text_from_docx() -> None:
    """Verify Word DOCX text extraction."""
    doc = docx.Document()
    doc.add_heading("Alice Developer", level=1)
    doc.add_paragraph("Senior Python & Distributed Systems Engineer")
    doc.add_paragraph("Architected real-time streaming pipelines.")

    docx_buffer = io.BytesIO()
    doc.save(docx_buffer)
    docx_bytes = docx_buffer.getvalue()

    extracted = extract_text_from_docx(docx_bytes)
    assert "Alice Developer" in extracted
    assert "Senior Python" in extracted

    parsed = parse_resume_bytes(docx_bytes, "Alice_Resume.docx")
    assert parsed.full_name == "Alice Developer"
    assert "Python" in parsed.skills


@pytest.mark.asyncio
async def test_candidate_onboarding_and_persona_api_flow(
    db_session: AsyncSession,
) -> None:
    """Verify complete onboarding API flow, persona switcher, and block CRUD."""

    from sqlalchemy.ext.asyncio import async_sessionmaker

    engine = db_session.bind
    session_maker = async_sessionmaker(
        bind=engine,
        class_=AsyncSession,
        expire_on_commit=False,
    )

    async def override_get_db() -> AsyncGenerator[AsyncSession, None]:
        async with session_maker() as session:
            try:
                yield session
                await session.commit()
            except Exception:
                await session.rollback()
                raise

    app.dependency_overrides[get_db] = override_get_db
    try:
        transport = ASGITransport(app=app)
        async with AsyncClient(transport=transport, base_url="http://test") as client:
            token = create_access_token(
                user_id="11111111-2222-3333-4444-555555555555",
                email="cand100@example.com",
            )
            headers = {"Authorization": f"Bearer {token}"}

            # 1. Parse text endpoint
            paste_payload = {
                "text": (
                    "Bob Engineer\n"
                    "bob@example.com | (555) 019-2831 | github.com/bobeng\n\n"
                    "# Experience\n"
                    "Senior Backend Engineer at DataCo\n"
                    "Architected high-throughput ingestion in Python & FastAPI "
                    "processing 50k req/s with 35% latency reduction.\n\n"
                    "# Hackathons & Honors\n"
                    "1st Place Winner at Global AI Hackathon\n"
                    "Built autonomous multi-agent task orchestrator using LangGraph."
                )
            }
            parse_resp = await client.post(
                "/api/candidates/parse-text", json=paste_payload
            )
            assert parse_resp.status_code == 200
            staged = parse_resp.json()
            assert staged["full_name"] == "Bob Engineer"
            assert staged["email"] == "bob@example.com"
            assert len(staged["blocks"]) >= 2

            # 2. Onboard candidate with verified blocks & specialized personas
            onboard_payload = {
                "full_name": "Bob Engineer",
                "visa_status": "H-1B",
                "citizenship_country_code": "US",
                "preferred_locations": ["San Francisco, CA", "Remote"],
                "remote_preference": "Remote",
                "target_roles": ["Senior Backend Engineer", "AI Platform Engineer"],
                "target_industries": ["ai_ml_infra", "hft_quant_trading"],
                "years_experience": 6,
                "seniority_level": "senior",
                "skills": ["Python", "FastAPI", "Kafka", "Docker", "PostgreSQL"],
                "blocks": staged["blocks"],
                "specialized_personas": [
                    {
                        "name": "AI Platform Engineer",
                        "headline": "Specializing in Large-Scale AI Orchestration",
                        "target_roles": ["AI Platform Engineer"],
                        "target_industries": ["ai_ml_infra"],
                        "seniority_level": "senior",
                        "years_experience": 4,
                        "skills": ["Python", "LangGraph", "FastAPI"],
                    }
                ],
            }
            onboard_resp = await client.post(
                "/api/candidates/onboard",
                headers=headers,
                json=onboard_payload,
            )
            assert onboard_resp.status_code == 200
            onboard_data = onboard_resp.json()
            assert onboard_data["status"] == "onboarded"
            assert onboard_data["blocks_ingested"] >= 2
            assert onboard_data["specialized_personas_created"] == 1

            # 3. List candidate personas
            personas_resp = await client.get(
                "/api/candidates/personas", headers=headers
            )
            assert personas_resp.status_code == 200
            personas_list = personas_resp.json()
            # Must contain at least Universal Persona + Specialized Persona
            assert len(personas_list) >= 2
            universal = next((p for p in personas_list if p["is_common_persona"]), None)
            specialized = next(
                (p for p in personas_list if not p["is_common_persona"]), None
            )
            assert universal is not None
            assert specialized is not None
            assert universal["is_default"] is True
            assert specialized["name"] == "AI Platform Engineer"

            # 4. Set specialized persona as active default
            set_def_resp = await client.put(
                f"/api/candidates/personas/{specialized['id']}/default",
                headers=headers,
            )
            assert set_def_resp.status_code == 200

            # 5. List CV blocks
            blocks_resp = await client.get("/api/candidates/blocks", headers=headers)
            assert blocks_resp.status_code == 200
            blocks_list = blocks_resp.json()
            assert len(blocks_list) >= 2
            assert any(b["has_embedding"] for b in blocks_list)

            # 6. Add new hackathon achievement block
            new_block_payload = {
                "category": "hackathon_achievement",
                "title": "1st Place Winner - AI Infra Track",
                "organization": "ETHGlobal",
                "content": (
                    "Built open-source decentralized LLM scheduler in Go & Kafka."
                ),
                "metrics": {"placement": "1st Place", "github_stars": 350},
                "skills": ["Go", "Kafka", "Distributed Systems"],
            }
            add_block_resp = await client.post(
                "/api/candidates/blocks",
                headers=headers,
                json=new_block_payload,
            )
            assert add_block_resp.status_code == 201
            new_block_id = add_block_resp.json()["id"]

            # 7. Delete block
            del_resp = await client.delete(
                f"/api/candidates/blocks/{new_block_id}",
                headers=headers,
            )
            assert del_resp.status_code == 200
            assert del_resp.json()["status"] == "deleted"
    finally:
        app.dependency_overrides.pop(get_db, None)


@pytest.mark.asyncio
async def test_persona_management_and_uniqueness_flow(
    db_session: AsyncSession,
) -> None:
    """
    Verify persona editing, renaming, case-insensitive duplicate guards,
    and Universal Profile rename/delete protections.
    """
    from sqlalchemy.ext.asyncio import async_sessionmaker

    engine = db_session.bind
    session_maker = async_sessionmaker(
        bind=engine,
        class_=AsyncSession,
        expire_on_commit=False,
    )

    async def override_get_db() -> AsyncGenerator[AsyncSession, None]:
        async with session_maker() as session:
            try:
                yield session
                await session.commit()
            except Exception:
                await session.rollback()
                raise

    app.dependency_overrides[get_db] = override_get_db
    try:
        transport = ASGITransport(app=app)
        async with AsyncClient(transport=transport, base_url="http://test") as client:
            user_id = "22222222-3333-4444-5555-666666666666"
            token = create_access_token(
                user_id=user_id,
                email="persona_mgr@example.com",
            )
            headers = {"Authorization": f"Bearer {token}"}

            # 1. Onboard candidate to create root account and Universal Profile
            onboard_payload = {
                "full_name": "Persona Manager Test",
                "visa_status": "Citizen",
                "target_roles": ["Software Engineer"],
                "years_experience": 5,
                "skills": ["Python", "FastAPI"],
            }
            onb_res = await client.post(
                "/api/candidates/onboard",
                headers=headers,
                json=onboard_payload,
            )
            assert onb_res.status_code == 200

            # 2. Retrieve personas (initially contains Universal Profile)
            list_res = await client.get("/api/candidates/personas", headers=headers)
            assert list_res.status_code == 200
            personas = list_res.json()
            assert len(personas) == 1
            universal_p = personas[0]
            assert universal_p["is_common_persona"] is True
            assert universal_p["name"] == "Universal Profile"

            # 3. Create a specialized persona "AI Engineer"
            create_payload = {
                "name": "AI Engineer",
                "headline": "LLM Infrastructure Specialist",
                "target_roles": ["AI Engineer"],
                "seniority_level": "senior",
                "years_experience": 4,
                "skills": ["PyTorch", "vLLM"],
            }
            create_res = await client.post(
                "/api/candidates/personas",
                headers=headers,
                json=create_payload,
            )
            assert create_res.status_code == 201
            ai_persona = create_res.json()
            assert ai_persona["name"] == "AI Engineer"

            # 4. Attempt to create duplicate name (case-insensitive "ai engineer")
            dup_create_res = await client.post(
                "/api/candidates/personas",
                headers=headers,
                json={"name": "ai engineer", "skills": ["Python"]},
            )
            assert dup_create_res.status_code == 400
            assert "already exists" in dup_create_res.json()["detail"]

            # 5. Create a second specialized persona "Backend Specialist"
            create_res2 = await client.post(
                "/api/candidates/personas",
                headers=headers,
                json={"name": "Backend Specialist", "skills": ["Go", "Kafka"]},
            )
            assert create_res2.status_code == 201
            backend_persona = create_res2.json()

            # 6. Rename "Backend Specialist" to "Distributed Systems Architect"
            update_payload = {
                "name": "Distributed Systems Architect",
                "headline": "High-Throughput Systems",
                "target_roles": ["Principal Architect"],
                "seniority_level": "principal",
                "years_experience": 10,
                "skills": ["Go", "Kafka", "Rust"],
            }
            update_res = await client.put(
                f"/api/candidates/personas/{backend_persona['id']}",
                headers=headers,
                json=update_payload,
            )
            assert update_res.status_code == 200
            updated_p = update_res.json()
            assert updated_p["name"] == "Distributed Systems Architect"
            assert updated_p["years_experience"] == 10

            # 7. Attempt to rename to "ai engineer" (conflict with Persona 1)
            conflict_update_res = await client.put(
                f"/api/candidates/personas/{backend_persona['id']}",
                headers=headers,
                json={
                    "name": "ai engineer",
                    "headline": "Conflicting",
                    "target_roles": [],
                },
            )
            assert conflict_update_res.status_code == 400
            assert "already exists" in conflict_update_res.json()["detail"]

            # 8. Attempt to rename the Universal Profile (must be rejected)
            universal_rename_res = await client.put(
                f"/api/candidates/personas/{universal_p['id']}",
                headers=headers,
                json={
                    "name": "Custom Profile Name",
                    "headline": "Changed headline",
                    "target_roles": ["CTO"],
                },
            )
            assert universal_rename_res.status_code == 400
            assert "cannot be renamed" in universal_rename_res.json()["detail"]

            # 9. Modifying settings of Universal Profile with same name succeeds
            universal_update_res = await client.put(
                f"/api/candidates/personas/{universal_p['id']}",
                headers=headers,
                json={
                    "name": "Universal Profile",
                    "headline": "Updated Universal Headline",
                    "target_roles": ["Director of Engineering"],
                    "seniority_level": "lead",
                    "years_experience": 8,
                    "skills": ["Python", "FastAPI", "PostgreSQL"],
                },
            )
            assert universal_update_res.status_code == 200
            upd_univ = universal_update_res.json()
            assert upd_univ["headline"] == "Updated Universal Headline"
            assert upd_univ["seniority_level"] == "lead"

            # 10. Attempt to delete Universal Profile (must be rejected)
            del_univ_res = await client.delete(
                f"/api/candidates/personas/{universal_p['id']}",
                headers=headers,
            )
            assert del_univ_res.status_code == 400
            assert "cannot be deleted" in del_univ_res.json()["detail"]

            # 11. Set specialized persona as default, then delete it
            await client.put(
                f"/api/candidates/personas/{ai_persona['id']}/default",
                headers=headers,
            )
            del_ai_res = await client.delete(
                f"/api/candidates/personas/{ai_persona['id']}",
                headers=headers,
            )
            assert del_ai_res.status_code == 200
            assert del_ai_res.json()["status"] == "deleted"

            # 12. Verify Universal Profile is automatically restored as default
            list_res2 = await client.get("/api/candidates/personas", headers=headers)
            personas2 = list_res2.json()
            univ_after = next(p for p in personas2 if p["is_common_persona"])
            assert univ_after["is_default"] is True
    finally:
        app.dependency_overrides.pop(get_db, None)
